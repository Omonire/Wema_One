#!/usr/bin/env python3
"""Rough text-overflow audit for the Luma deck: flags shapes whose text
likely exceeds the shape bounds, and shapes sticking out of the slide."""

import os
import sys

from pptx import Presentation
from pptx.util import Pt

PATH = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else "Luma_Pitch_Deck.pptx")
EMU_PER_PT = 12700
SLIDE_W_PT = 13.333 * 72
SLIDE_H_PT = 7.5 * 72

# Average glyph width as a fraction of font size (rough, per font family class)
WIDTH_FACTOR = 0.52


def para_height(p, width_pt):
    if not p.runs:
        return 0.0
    size = max((r.font.size.pt if r.font.size else 18) for r in p.runs)
    text = "".join(r.text for r in p.runs)
    # estimate wrapped line count
    text_w = len(text) * size * WIDTH_FACTOR
    lines = max(1, math_ceil(text_w / max(width_pt, 1)))
    line_h = size * 1.22
    return lines * line_h + (p.space_before.pt if p.space_before else 0) + (p.space_after.pt if p.space_after else 0)


def math_ceil(x):
    import math
    return math.ceil(x)


def main():
    prs = Presentation(PATH)
    issues = 0
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            name = shape.shape_type
            # off-slide check
            try:
                l, t = shape.left, shape.top
                w, h = shape.width, shape.height
            except Exception:
                continue
            if l is None:
                continue
            if l < -10000 or t < -10000 or (l + w) > int(SLIDE_W_PT * EMU_PER_PT) + 10000 or (t + h) > int(SLIDE_H_PT * EMU_PER_PT) + 10000:
                # allow decorative blobs that intentionally bleed
                if shape.name.startswith(("Oval", "Rectangle")) and (l < 0 or t < 0 or l + w > SLIDE_W_PT * EMU_PER_PT or t + h > SLIDE_H_PT * EMU_PER_PT):
                    continue  # intentional bleed decoration
                print(f"[S{idx:02d}] OFF-SLIDE: {name} at ({l/914400:.2f}in, {t/914400:.2f}in) size {(w/914400):.2f}x{(h/914400):.2f}in")
                issues += 1
            if not shape.has_text_frame:
                continue
            tf = shape.text_frame
            text = tf.text.strip()
            if not text:
                continue
            w_pt = (w / EMU_PER_PT) - 4
            total = 0.0
            for p in tf.paragraphs:
                total += para_height(p, w_pt)
            h_pt = h / EMU_PER_PT
            if total > h_pt * 1.06:  # 6% tolerance
                print(f"[S{idx:02d}] TEXT OVERFLOW ~{total:.0f}pt in {h_pt:.0f}pt box: "
                      f"\"{text[:60].replace(chr(10), ' / ')}...\"")
                issues += 1
    print(f"\nDone. {issues} potential issue(s) found.")


if __name__ == "__main__":
    main()
